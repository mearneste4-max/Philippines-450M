# -*- coding: utf-8 -*-
"""菲律宾 450MHz 方案 v7

五段式骨架（Why 450MHz & Hytera → Industry solutions → E2E architecture →
Reference cases → Business model），浅色配色：白底 + 海能达蓝标题 + 橄榄绿点缀
+ 浅灰蓝卡片。

用法：python build_ph450m_v7.py [模板.pptx] [输出.pptx]
  模板缺省为 01_交付物/菲律宾450M方案_v7.pptx（脚本会清空它的全部页面，只复用
  母版、版式与 notesMaster）；输出缺省为原地覆盖模板文件（旧版本靠 git 追溯）。
  素材优先在仓库内查找（05_素材库、06_用到的图片、03_生成脚本/素材），找不到再查
  F 盘 Workspace 下的原始路径，因此仓库可独立重建。
"""
import copy
import hashlib
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR as ANCHOR
from pptx.enum.text import PP_ALIGN as A
from pptx.oxml.ns import qn
from pptx.util import Pt

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
TPL = sys.argv[1] if len(sys.argv) > 1 else str(REPO / '01_交付物' / '菲律宾450M方案_v7.pptx')
OUT = sys.argv[2] if len(sys.argv) > 2 else TPL   # 缺省原地重建（模板已整体读入内存）
G = HERE / '素材'
(G / '_hi').mkdir(parents=True, exist_ok=True)   # 预处理缓存，不入库

# ---------------------------------------------------------------- 设计令牌
C = {
    'WHITE': 'FFFFFF', 'PAPER': 'F7FAFD', 'CARD': 'EEF2F9', 'CARD2': 'E4EBF7',
    'LINE': 'D6E0F0', 'NAVY': '0B308E', 'NAVY_L': '3A5BA8', 'NAVY_D': '07225F',
    'BLUE_M': '8BB0CD', 'GREEN': '7F9B00', 'LIME': '8CC63F', 'GREEN_L': 'F1F6E2',
    'ORANGE': 'E8912A', 'INK': '16213E', 'STEEL': '55607A', 'SILVER': '93A0B4',
    'MIST': 'DCE5F3',
}
W, H = 960, 540
X0, X1 = 48, 912
CW = X1 - X0
EN = 'Arial'
L, CEN, R = A.LEFT, A.CENTER, A.RIGHT
TOP, MID = ANCHOR.TOP, ANCHOR.MIDDLE

# 仓库内素材排在最前，保证同名文件优先取仓库版本
ROOTS = [G / '_recovered', REPO / '06_用到的图片', REPO / '05_素材库']
ROOTS += [G / n for n in ('_orig', '_case', '_mcs', '_prod', '_cut', '_hi', '_ex', 'images')]
ROOTS += [G / '_casepool' / n / 'images' for n in
          ('PNG', 'ICCS', 'HYTALK', 'MCX', 'TETRA2', 'SDP', 'MBB', 'RAN', 'QOS')]
ROOTS += [Path(r'F:\Workspace\03_sources\references\2026-10_资源中心图库\Approved_可公开使用'),
          Path(r'F:\Workspace\04_templates\素材库'),
          Path(r'F:\Workspace\03_sources\references\2026-09_应急快速部署\素材图片'),
          Path(r'F:\Workspace\01_projects\2026-10_应急产品内部分享\working\_gen\素材\images_v3')]
_IX = {}
for r in ROOTS:
    if r.is_dir():
        for p in r.rglob('*'):
            if p.suffix.lower() in ('.png', '.jpg', '.jpeg'):
                _IX.setdefault(p.name, p)


def asset(name):
    p = Path(name)
    if p.is_absolute() and p.exists():
        return p
    if Path(name).name in _IX:
        return _IX[Path(name).name]
    raise FileNotFoundError(name)


_IMG = {}


def slim(src, maxw=1600):
    """嵌入前统一缩放并转 JPEG，避免 PPTX 体积失控。"""
    src = Path(src)
    key = hashlib.md5(str(src).encode()).hexdigest()[:10]
    out = G / '_hi' / ('v6_%s.jpg' % key)
    if not out.exists():
        from PIL import Image
        im = Image.open(src)
        if im.mode in ('RGBA', 'LA', 'P'):
            im = im.convert('RGBA')
            bg = Image.new('RGB', im.size, (255, 255, 255))
            bg.paste(im, mask=im.split()[-1])
            im = bg
        else:
            im = im.convert('RGB')
        if im.width > maxw:
            im = im.resize((maxw, max(1, int(im.height * maxw / im.width))), Image.LANCZOS)
        im.save(out, quality=88, optimize=True)
    return out


def cropcache(name, src, box):
    """按比例裁掉不要的区域（如报告页边框），缓存复用。"""
    out = G / '_hi' / name
    if not out.exists():
        from PIL import Image
        im = Image.open(src).convert('RGB')
        w, h = im.size
        im.crop((int(w * box[0]), int(h * box[1]), int(w * box[2]), int(h * box[3]))).save(out)
    return out


def illu(name, maxw=1200):
    """等距插画：保留透明通道，缩到 maxw 以内后存 PNG。"""
    p = asset(name)
    out = G / '_hi' / ('i_%s.png' % p.stem)
    if not out.exists():
        from PIL import Image
        im = Image.open(p).convert('RGBA')
        if im.width > maxw:
            im = im.resize((maxw, max(1, int(im.height * maxw / im.width))), Image.LANCZOS)
        im.save(out)
    return out


def enhance(name, src, minw=1500):
    """放大 + 锐化。src 可传文件名（延迟解析）：仓库里已有同名成品时直接用成品，
    原始素材不在本机也能重建。"""
    from PIL import Image, ImageFilter
    out = G / '_hi' / name
    if not out.exists() and name in _IX:
        return _IX[name]
    if not out.exists():
        im = Image.open(asset(src) if isinstance(src, str) else src).convert('RGB')
        if im.width < minw:
            im = im.resize((minw, int(im.height * minw / im.width)), Image.LANCZOS)
        im.filter(ImageFilter.UnsharpMask(radius=2, percent=110, threshold=3)).save(out)
    return out


# ---------------------------------------------------------------- 基础图元
def box(s, x, y, w, h, fill=None, outline=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE,
        shadow=False):
    sh = s.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sh.fill.background()
    if outline:
        sh.line.color.rgb = RGBColor.from_string(outline); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    if not shadow:
        sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def ln(s, x1, y1, x2, y2, color, w=1.0, dash=False):
    c = s.shapes.add_connector(1, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    c.line.color.rgb = RGBColor.from_string(color); c.line.width = Pt(w)
    if dash:
        c.line.dash_style = 4
    return c


def tx(s, x, y, w, h, text, size=11, bold=False, color='INK', align=L, anchor=TOP,
       line=1.25, after=3, font=EN):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h)); tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, seg in enumerate(str(text).split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        p.space_after = Pt(after)
        r = p.add_run(); r.text = seg
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = font
        r.font.color.rgb = RGBColor.from_string(C.get(color, color))
    return tb


def rich(s, x, y, w, h, parts, size=11, anchor=TOP, line=1.3, align=L):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h)); tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align; p.line_spacing = line
    for t, b, col in parts:
        r = p.add_run(); r.text = t
        r.font.size = Pt(size); r.font.bold = b; r.font.name = EN
        r.font.color.rgb = RGBColor.from_string(C.get(col, col))
    return tb


def keepalpha(src, maxw=900):
    """PNG（图标／插画）只做缩放，保留透明通道——绝不能走 slim() 的白底合成。"""
    src = Path(src)
    if src.suffix.lower() != '.png':
        return slim(src)
    key = hashlib.md5(str(src).encode()).hexdigest()[:10]
    out = G / '_hi' / ('a_%s.png' % key)
    if not out.exists():
        from PIL import Image
        im = Image.open(src)
        if im.mode not in ('RGBA', 'LA', 'P'):
            im = im.convert('RGBA')
        if im.width > maxw:
            im = im.resize((maxw, max(1, int(im.height * maxw / im.width))), Image.LANCZOS)
        im.save(out, optimize=True)
    return out


def fit(s, src, x, y, w, h):
    from PIL import Image
    src = keepalpha(src)
    with Image.open(src) as im:
        iw, ih = im.size
    r = min(w / iw, h / ih)
    return s.shapes.add_picture(str(src), Pt(x + (w - iw * r) / 2),
                                Pt(y + (h - ih * r) / 2), Pt(iw * r), Pt(ih * r))


def photo(s, src, x, y, w, h, fy=0.5, fx=0.5):
    from PIL import Image
    src = slim(src)
    with Image.open(src) as im:
        iw, ih = im.size
    p = s.shapes.add_picture(str(src), Pt(x), Pt(y), Pt(w), Pt(h))
    tgt, sr = w / h, iw / ih
    if sr > tgt:
        keep = tgt / sr; left = min(max(fx - keep / 2, 0), 1 - keep)
        p.crop_left, p.crop_right = left, 1 - keep - left
    else:
        keep = sr / tgt; top = min(max(fy - keep / 2, 0), 1 - keep)
        p.crop_top, p.crop_bottom = top, 1 - keep - top
    return p


# ---------------------------------------------------------------- 页面家具
def page(s, n, kicker, title, point=None):
    """浅色页头：白底 + 深蓝标题 + 右侧章节标签与页码 + 绿色细线。"""
    box(s, 0, 0, W, H, C['WHITE'])
    box(s, X0, 30, 18, 5, C['LIME'])
    tx(s, X0 + 26, 22, 780, 30, title, 21, True, 'NAVY')
    tx(s, X0, 62, 700, 16, kicker.upper(), 8.5, True, 'GREEN')
    tx(s, X1 - 44, 22, 44, 26, '%02d' % n, 17, True, 'NAVY', R)
    ln(s, X0, 82, X1, 82, C['LINE'], 1.0)
    if point:
        tx(s, X0, 94, CW, 18, point, 11.5, False, 'STEEL')


def concl(s, parts, y=452):
    box(s, 0, y, W, 40, C['NAVY'])
    rich(s, X0, y, CW, 40, parts, 12, anchor=MID)


def src(s, text):
    tx(s, X0, 508, CW, 14, 'Source: ' + text, 8, False, 'SILVER')


def stat(s, x, y, w, h, value, label, accent='NAVY', vsize=22):
    box(s, x, y, w, h, C['CARD'])
    box(s, x, y, 3.5, h, C[accent])
    tx(s, x + 16, y + 12, w - 32, 28, value, vsize, True, accent)
    tx(s, x + 16, y + 44, w - 32, h - 50, label, 9.5, False, 'STEEL', line=1.25)


def klab(s, x, y, text, color='GREEN'):
    tx(s, x, y, min(420, X1 - x), 14, text.upper(), 8.5, True, color)


def bullets(s, x, y, w, items, size=11, gap=34, dot='LIME'):
    for i, t in enumerate(items):
        yy = y + i * gap
        box(s, x, yy + 5, 6, 6, C[dot])
        tx(s, x + 16, yy, w - 16, gap - 4, t, size, False, 'STEEL', line=1.28)


# ---------------------------------------------------------------- 模板准备
prs = Presentation(TPL)
while len(prs.slides._sldIdLst):
    r = prs.slides._sldIdLst[0]
    prs.part.drop_rel(r.rId)
    prs.slides._sldIdLst.remove(r)
BLANK = None
for lay in prs.slide_layouts:
    if lay.name.lower() in ('blank', '空白'):
        BLANK = lay
BLANK = BLANK or prs.slide_layouts[6]


def newslide():
    return prs.slides.add_slide(BLANK)


# ================================================================ 01 封面
s = newslide()
box(s, 0, 0, W, H, C['WHITE'])
box(s, 0, 0, W, 6, C['NAVY'])
box(s, 0, 0, 320, 6, C['LIME'])
box(s, 0, 6, 6, H - 6, C['NAVY'])

# 右上角品牌标识
fit(s, illu('hytera_logo.png', 900), X1 - 148, 30, 110, 44)

# 左：标题区
tx(s, X0 + 20, 108, 460, 18, 'HYTERA  \u00d7  GLOBE TELECOM', 10.5, True, 'GREEN')
tx(s, X0 + 20, 138, 470, 120, '450 MHz Mission-Critical\nNetwork for the\nPhilippines',
   34, True, 'NAVY', line=1.14)
box(s, X0 + 20, 296, 72, 4, C['LIME'])
tx(s, X0 + 20, 316, 440, 60,
   'A coverage and resilience layer for utilities,\npublic safety, transport and remote communities.',
   12.5, False, 'STEEL', line=1.42)

# 右：单一焦点图——菲律宾地图 + 终端实物
fit(s, illu('\u83f2\u5f8b\u5bbe\u89e3\u51b3\u65b9\u6848\u5305\u88c5_slide11_1.png', 1400), 628, 66, 244, 352)
fit(s, illu('\u7ec8\u7aef_\u591a\u6a21\u624b\u53f0_1.png'), 542, 148, 86, 266)

# 底部：一条细线与两行小字
ln(s, X0, 466, X1, 466, C['LINE'], 1.0)
tx(s, X0, 482, 400, 16, 'Philippines  \u00b7  October 2026', 11.5, True, 'NAVY')
tx(s, 560, 484, 352, 14, 'Terminal  \u00b7  Network  \u00b7  Application', 9.5, False, 'SILVER', R)

# ================================================================ 02 国情
s = newslide()
page(s, 2, 'The gap', 'Islands, mountains and typhoons',
     'Public 4G and 5G reach the cities. What is missing is coverage that is still there after the storm.')
box(s, X0, 122, 400, 296, C['PAPER'])
fit(s, illu('\u83f2\u5f8b\u5bbe\u89e3\u51b3\u65b9\u6848\u5305\u88c5_slide11_1.png', 1400), X0 + 22, 140, 150, 260)
box(s, X0 + 186, 140, 196, 96, C['CARD'])
tx(s, X0 + 202, 154, 166, 70, '7,641 islands.\nOne grid has to reach\nacross all of them.',
   10.5, True, 'NAVY', line=1.35)
box(s, X0 + 186, 248, 196, 78, C['WHITE'])
fit(s, illu('\u5c71\u8109\u7ebf\u7a3f.png'), X0 + 194, 256, 180, 56)
fit(s, illu('\u901a\u4fe1\u94c1\u5854_3.png'), X0 + 316, 258, 34, 62)
box(s, X0 + 186, 334, 196, 72, C['CARD'])
tx(s, X0 + 200, 344, 168, 56,
   'Mountains, storm landfall and\n2,000 km of sea between\nthe islands.',
   9.5, False, 'STEEL', line=1.35)
STATS = [('7,641', 'islands, much of the terrain mountainous', 'NAVY'),
         ('#1', 'WorldRiskIndex 2025 - highest disaster risk', 'GREEN'),
         ('~20', 'tropical cyclones enter PAR every year', 'NAVY'),
         ('1.4M', 'electricity consumers without power after Typhoon Tino (Nov 2025)', 'GREEN')]
for i2, (v, lab, ac) in enumerate(STATS):
    r, c = divmod(i2, 2)
    stat(s, 480 + c * 220, 122 + r * 152, 208, 140, v, lab, ac, 24)
concl(s, [('450 MHz is the layer that stays up.  ', True, 'WHITE'),
          ('Coverage and resilience, not capacity - alongside Globe\u2019s 4G and 5G, not instead of them.',
           False, 'MIST')])
src(s, 'WorldRiskReport 2025; PAGASA; NAMRIA; Department of Energy, Typhoon Tino power situation (5 Nov 2025)')

# ================================================================ 03 目录
s = newslide()
page(s, 3, 'Agenda', 'What we will cover',
     'Five sections. The first two are the argument, the last three are the delivery.')
AG = [('01', 'Why 450 MHz & Hytera', 'Global trend, the band itself, and what MCS means',
       'thumb_why450.png', 'P4 \u2013 P6'),
      ('02', 'Industry solutions', 'Power \u00b7 Water \u00b7 Police \u00b7 Emergency \u00b7 Rail \u00b7 Port',
       'thumb_industries.png', 'P7 \u2013 P13'),
      ('03', 'End-to-end architecture', 'Terminal \u00b7 Network \u00b7 Application',
       'thumb_architecture.png', 'P14'),
      ('04', 'Reference cases', 'Global 450 MHz networks, and three customer cases',
       'thumb_cases.png', 'P15 \u2013 P18'),
      ('05', 'Business model', 'Who does what, six revenue streams, next steps',
       'thumb_model.png', 'P19 \u2013 P21')]
for i, (num, title, sub, thumb, pgs) in enumerate(AG):
    x = X0 + i * 174
    box(s, x, 122, 162, 296, C['CARD'])
    fit(s, illu(thumb, 396), x, 122, 162, 82)
    ln(s, x, 204, x + 162, 204, C['LINE'], 1.0)
    tx(s, x + 14, 214, 134, 30, num, 21, True, 'LIME')
    tx(s, x + 14, 250, 134, 40, title, 11.5, True, 'NAVY', line=1.2)
    tx(s, x + 14, 296, 134, 66, sub, 9, False, 'STEEL', line=1.3)
    tx(s, x + 14, 388, 134, 14, pgs, 8.5, True, 'SILVER')
concl(s, [('One network, six industries.  ', True, 'WHITE'),
          ('The same 450 MHz layer serves all of them, on the same sites.', False, 'MIST')])
src(s, 'Hytera proposal structure')

# ================================================================ 04 450MHz 全球
s = newslide()
page(s, 4, 'Global trend', 'From a niche band to critical infrastructure',
     'Regulators, utilities and public safety agencies have settled on what this band is for.')
TREND = [('86%', 'of the world\u2019s population lives in countries where the 400 MHz bands are allocated or under consultation', 'NAVY'),
         ('77 / 39', 'countries with networks deployed at 380 / 410 / 450 MHz - and the 39 of them already running LTE', 'GREEN'),
         ('1.8M / 14M', 'broadband and sensor connections forecast by 2030, from 420k and 1.2M in 2026', 'NAVY')]
for i, (v, lab, ac) in enumerate(TREND):
    stat(s, X0, 122 + i * 86, 470, 78, v, lab, ac, 21)
POS = [('Named at both ends of the chain',
        'The Alliance lists Hytera in its device supplier list and in its macro eNB supplier list.'),
       ('Through operator acceptance',
        'PNC660 450M is whitelisted by 450connect to \u2212115 dBm RSRP - the operator\u2019s own document.'),
       ('A 5G-ready 450 MHz portfolio',
        'Handheld, vehicle radio, base station and core, with a 5G model in development.')]
klab(s, 546, 122, 'Our position')
for i, (t, d) in enumerate(POS):
    y = 142 + i * 84
    box(s, 546, y, 366, 76, C['CARD'])
    box(s, 546, y, 3.5, 76, C['NAVY'] if i == 0 else C['MIST'])
    tx(s, 564, y + 12, 336, 18, t, 11.5, True, 'NAVY')
    tx(s, 564, y + 34, 336, 40, d, 9, False, 'STEEL', line=1.3)
concl(s, [('The band is no longer experimental.  ', True, 'WHITE'),
          ('It is assigned, standardised in 3GPP, and carrying utility and public-safety traffic today.',
           False, 'MIST')])
src(s, '450 MHz Alliance, Annual Global Update (Dec 2025) p4, p11 Table 1, p12 Table 6; Annual Equipment Report (May 2025) p5')

# ================================================================ 05 物理层面
s = newslide()
page(s, 5, 'Why 450 MHz', 'One site replaces several mid-band sites',
     'Low-band physics is the whole argument. The numbers below reuse one grid at three frequencies.')
DATA = [('410 / 450 MHz', '1\u00d7 sites  \u00b7  100% covered', [(0, 0)], 62),
        ('740 MHz', '2.5-3\u00d7 sites  \u00b7  91.3% from the same grid',
         [(-36, -16), (36, -16), (0, 28)], 34),
        ('2350 MHz', '12-15\u00d7 sites  \u00b7  58.7% from the same grid',
         [(i * 38 - 57, j * 30 - 45) for i in range(4) for j in range(4)
          if (i, j) not in ((0, 0), (3, 0), (0, 3))], 17)]
for i, (name, sub, sites, r) in enumerate(DATA):
    x = X0 + i * 166
    box(s, x, 122, 156, 156, C['PAPER'])
    for k in range(1, 4):
        ln(s, x + k * 39, 127, x + k * 39, 273, C['MIST'], 0.5)
        ln(s, x + 5, 122 + k * 39, x + 151, 122 + k * 39, C['MIST'], 0.5)
    cx, cy = x + 78, 200
    for dx, dy in sites:
        box(s, cx + dx - r, cy + dy - r * 0.86, 2 * r, 2 * r * 0.86,
            'D8E4F6', C['NAVY'], 0.75, shape=MSO_SHAPE.OVAL)
    tx(s, x, 288, 156, 16, name, 11, True, 'NAVY', CEN)
    tx(s, x, 306, 156, 30, sub, 9, True, 'GREEN', CEN, line=1.25)
tx(s, X0, 344, 500, 32,
   'Same area, same grid. Only the low band covers it with one site - which is what makes island and\n'
   'mountain coverage affordable, and what makes extensive backup power affordable.',
   10, False, 'STEEL', line=1.35)
klab(s, 560, 122, 'What that buys')
bullets(s, 560, 148, 352, [
    'Bends around ridges and into valleys instead of needing line of sight',
    'Reaches meter boxes, basements and tunnels that mid-band cannot',
    'Fewer sites is what makes extensive backup power affordable',
], 11.5, 50)
concl(s, [('What it is not:  ', True, 'WHITE'),
          ('with 3-5 MHz it carries voice, low-rate video and telemetry. It stays outside the capacity game.',
           False, 'MIST')])
src(s, '450 MHz Alliance, Brief (Jun 2022) p15 - reference grid modelled at 410 MHz')

# ================================================================ 06 MCS
s = newslide()
page(s, 6, 'The service layer', 'What MCS means - and how it sits on the band',
     'The customer cases in section 04 run MCS. This page explains what it is, because the term carries the delivery.')
box(s, X0, 122, 430, 200, C['CARD'])
tx(s, X0 + 18, 136, 394, 18, '3GPP mission-critical services (MCS / MCX)', 12, True, 'NAVY')
MCS = [('\u7ebf\u7a3f_\u624b\u53f0_2.png', 'MCPTT', 'push-to-talk voice, with priority and pre-emption'),
       ('\u7ebf\u7a3f_\u624b\u53f0_3.png', 'MCVideo', 'live and recorded video from the scene'),
       ('\u7ebf\u7a3f_\u7ec8\u7aef_1.png', 'MCData', 'messaging, file transfer and location data')]
for i, (ic, k, d) in enumerate(MCS):
    y = 164 + i * 48
    box(s, X0 + 18, y, 394, 40, C['WHITE'])
    box(s, X0 + 18, y, 3, 40, C['LIME'])
    fit(s, illu(ic), X0 + 30, y + 6, 16, 28)
    tx(s, X0 + 54, y + 4, 110, 18, k, 11.5, True, 'NAVY')
    tx(s, X0 + 54, y + 21, 344, 16, d, 9, False, 'STEEL')
tx(s, X0, 330, 430, 44,
   'Defined by 3GPP, not by a single vendor - so the network stays multi-vendor and the\n'
   'terminals are not locked to one supplier.', 10, False, 'STEEL', line=1.35)
STACK = [('Applications', 'dispatch, recording, device management', 'NAVY'),
         ('MCS service layer', 'MCPTT / MCVideo / MCData', 'GREEN'),
         ('LTE or 5G radio', 'Band 31 / 72 at 450 MHz', 'NAVY_L')]
klab(s, 510, 122, 'How the layers stack')
for i, (t, d, col) in enumerate(STACK):
    y = 146 + i * 62
    box(s, 510, y, 402, 52, C[col])
    tx(s, 528, y + 8, 366, 18, t, 12, True, 'WHITE')
    tx(s, 528, y + 28, 366, 16, d, 9, False, 'MIST')
    if i < 2:
        box(s, 700, y + 54, 22, 16, C['NAVY_L'], shape=MSO_SHAPE.DOWN_ARROW)
box(s, 510, 336, 402, 38, C['GREEN_L'])
tx(s, 528, 336, 366, 38,
   'In this deck: the band carries the traffic, MCS provides the service.',
   10, True, 'GREEN', L, MID)
concl(s, [('450 MHz gives the coverage; MCS gives the service.  ', True, 'WHITE'),
          ('MCS runs on any 3GPP LTE or 5G radio, so it carries over to a 450 MHz layer unchanged.', False, 'MIST')])
src(s, '3GPP mission-critical services; Hytera HyTalk MC platform documentation')

# ================================================================ 07 六行业总览
s = newslide()
page(s, 7, 'Industry solutions', 'One network, six industries',
     'The same 450 MHz layer serves all six, on the same sites. Each of the next pages takes one.')
IND = [('Power', 'Grid dispatch, feeder automation and smart metering', '\u7535\u529b\u62a2\u4fee.png'),
       ('Water', 'Metering, pressure sensing and reservoir telemetry', '\u6c34\u5e93\u5927\u575d.png'),
       ('Police', 'Cross-agency group calls and live video from the scene', '\u8b66\u5bdf\u5c40.png'),
       ('Emergency', 'Rapid deployment and coverage after the storm', '\u6d88\u9632\u7ad9\u706b\u707e.png'),
       ('Rail', 'Corridor coverage and train-to-ground voice and data', '\u4ea4\u901a\u67a2\u7ebd.png'),
       ('Port', 'Quay dispatch, ship-to-shore and coast guard patrols', '\u8239\u8236.png')]
for i, (t, d, ic) in enumerate(IND):
    r, c = divmod(i, 3)
    x = X0 + c * 292
    y = 122 + r * 152
    box(s, x, y, 280, 140, C['CARD'])
    box(s, x, y, 3.5, 140, C['NAVY'] if r == 0 else C['LIME'])
    box(s, x + 14, y + 12, 116, 116, C['PAPER'])
    fit(s, illu(ic), x + 20, y + 18, 104, 104)
    tx(s, x + 144, y + 24, 124, 22, t, 14.5, True, 'NAVY')
    tx(s, x + 144, y + 54, 124, 72, d, 9.5, False, 'STEEL', line=1.35)
concl(s, [('Utilities anchor it commercially.  ', True, 'WHITE'),
          ('Public safety, emergency, rail and port add national value on the same sites.', False, 'MIST')])
src(s, '450 MHz Alliance (Dec 2025); Hytera industry solutions')

# ================================================================ 08-13 六个行业页
IND_PAGES = [
    dict(n=8, kicker='Industry 1 of 6  \u00b7  Power', title='Power: the grid, the crews and the meters',
         point='Electric cooperatives have spread-out assets, a 10-15 year horizon and a willingness to pay for reliability.',
         hero=asset('20260908_\u6d77\u672c\u4e8c\u90e8\u57f9\u8bad_HyTalk-MC\u65b9\u6848_\u8bfe\u4ef6_V6.1_slide41_1.png'), fy=0.5,
         accent=asset('\u83f2\u5f8b\u5bbe\u89e3\u51b3\u65b9\u6848\u5305\u88c5_slide3_5.jpg'),
         cap='Grid repair crews and line assets. Illustrative imagery.',
         pain='Mountain and island lines have no reliable signal, and feeder automation needs a private data path that does not depend on the public network.',
         sol=['Wide-area push-to-talk, video and lone-worker alarms along the lines',
              'A private data channel for SCADA, feeder automation and substation control',
              'Smart meters on the same sites, with deep indoor coverage'],
         ph='In the Philippines: 121 electric cooperatives under NEA.',
         srcn='National Electrification Administration; Hytera power industry solutions'),
    dict(n=9, kicker='Industry 2 of 6  \u00b7  Water', title='Water: locate the losses instead of estimating them',
         point='Non-revenue water is usually estimated, not located. Pump, well and reservoir sites sit exactly where mid-band cannot reach.',
         hero=enhance('water_pumps.png', '\u83f2\u5f8b\u5bbe\u89e3\u51b3\u65b9\u6848\u5305\u88c5_slide7_5.jpg'), fy=0.5,
         accent=asset('\u83f2\u5f8b\u5bbe\u89e3\u51b3\u65b9\u6848\u5305\u88c5_slide5_3.jpg'),
         cap='Reservoir and pumping assets. Illustrative imagery.',
         pain='Most water districts know their losses as a percentage, not as an address - because the meters and pressure sensors cannot report.',
         sol=['Meters and pressure sensors that locate losses instead of estimating them',
              'Pump, well and reservoir telemetry from remote and underground sites',
              'Backup-powered sites keep alarms flowing when the public network is down'],
         ph='In the Philippines: only 149 of 532 water districts kept losses within LWUA\u2019s 20% limit (2025).',
         srcn='LWUA non-revenue water data, end-2025 (reported Aug 2026); Hytera water industry solutions'),
    dict(n=10, kicker='Industry 3 of 6  \u00b7  Police', title='Police: one group call across the agencies',
         point='The requirement is not another private network. It is interoperability between the agencies that already have one.',
         hero=asset('20241104_\u8d44\u6e90\u4e2d\u5fc3_ICC\u65b9\u6848_\u573a\u666f\u7248_V2_slide10_1.png'), fy=0.42,
         accent=asset('P60_Snow_Public safety.jpg'),
         cap='Police on traffic duty, and a handheld in the field. Illustrative imagery.',
         pain='Different agencies cannot talk to each other, and the public network congests exactly at the moment it is needed most.',
         sol=['One group call across police, fire and civil defence',
              'Live video from the scene, with priority and pre-emption when the network congests',
              'Interworking with the DMR and TETRA networks already in service'],
         ph='Reference: Singapore has assigned 450 MHz band 31 to public safety, in service.',
         srcn='450 MHz Alliance, Annual Global Update (Dec 2025) p8; Hytera public-safety references'),
    dict(n=11, kicker='Industry 4 of 6  \u00b7  Emergency', title='Emergency: coverage that survives the storm',
         point='After a typhoon the public network fails exactly where the response is needed. This is the layer that has to survive.',
         hero=asset('China Yunnan Fire Rescue_Flood Rescue Operation_Girls and Cat_1.jpg'), fy=0.42,
         accent=asset('Epack200_Snow.jpg'),
         cap='Flood rescue, and a portable base station in the field. Illustrative imagery.',
         pain='Rapid response needs communications within hours, not days - and it needs them where there is no grid and no backhaul.',
         sol=['Portable base stations and satellite backhaul deployed within hours',
              'A field command post that reaches the provincial EOC and the national agency',
              'Sites that keep running on backup power when the grid does not'],
         ph='Reference: rapid deployment kits and container command posts in service.',
         srcn='Hytera emergency and rapid-deployment product documentation'),
    dict(n=12, kicker='Industry 5 of 6  \u00b7  Rail', title='Rail: continuous coverage along a narrow corridor',
         point='A rail corridor is a narrow strip, often in cuttings. It is one of the two shapes mid-band handles worst.',
         hero=enhance('v6_metro.png', '20260920_\u4ea7\u54c1\u6587\u6863_TETRA\u65b9\u6848\u4ecb\u7ecd_\u8bfe\u4ef6_\u82f1\u6587_slide41_3.png'), fy=0.45,
         accent=asset('PNC660_Snow_Train.jpg'),
         cap='An elevated rail line, and a handheld on board. Illustrative imagery.',
         pain='Coverage has to follow the track, not the population - so a network built for towns leaves the corridor uncovered.',
         sol=['Continuous coverage along the corridor, including cuttings and tunnels',
              'Train-to-ground voice, data and video for operations and security',
              'Runs on the operator\u2019s network rather than a separate private system'],
         ph='Reference: South Africa has assigned the band to Transnet for rail.',
         srcn='450 MHz Alliance, Annual Global Update 2025 p6; Hytera railway references'),
    dict(n=13, kicker='Industry 6 of 6  \u00b7  Port', title='Port: from the quay out to sea',
         point='Ports are dense and metallic, and coast guard patrols run far beyond mid-band reach. One network covers both.',
         hero=asset('case_hzmb_bridge.jpg'), fy=0.5,
         accent=asset('sea_boat.png'),
         cap='A sea crossing, and a rescue boat. Illustrative imagery.',
         pain='Quay operations, ship-to-shore traffic and patrol boats all need the same group call, across water and steel.',
         sol=['Port and terminal dispatch on one group call across the quay',
              'Ship-to-shore voice and data from coastal sites',
              'Coast guard patrols covered far out to sea, on the same network as the port'],
         ph='The same coastal sites also serve island and rural communities.',
         srcn='450 MHz Alliance (Dec 2025); Hytera transport and maritime references'),
]
for cfg in IND_PAGES:
    s = newslide()
    page(s, cfg['n'], cfg['kicker'], cfg['title'], cfg['point'])
    photo(s, cfg['hero'], X0, 122, 380, 230, fy=cfg.get('fy', 0.5))
    box(s, X0, 122, 380, 230, None, C['LINE'], 1.0)
    box(s, X0 + 292, 134, 76, 92, C['WHITE'])
    photo(s, cfg['accent'], X0 + 296, 138, 68, 84, fy=0.5)
    tx(s, X0, 356, 380, 30, cfg['cap'], 8.5, False, 'SILVER', line=1.3)
    box(s, 456, 122, 456, 78, C['GREEN_L'])
    box(s, 456, 122, 3.5, 78, C['GREEN'])
    tx(s, 474, 132, 424, 14, 'PAIN POINT', 8.5, True, 'GREEN')
    tx(s, 474, 150, 424, 46, cfg['pain'], 9.5, False, 'STEEL', line=1.3)
    klab(s, 456, 210, 'What 450 MHz does')
    for j, it in enumerate(cfg['sol']):
        yy = 232 + j * 40
        fit(s, illu('\u7ebf\u7a3f_\u7ec8\u7aef_1.png'), 456, yy + 1, 12, 20)
        tx(s, 478, yy, 434, 36, it, 10.5, False, 'STEEL', line=1.28)
    box(s, 456, 356, 456, 30, C['CARD'])
    tx(s, 472, 356, 424, 30, cfg['ph'], 9.5, True, 'NAVY', L, MID)
    src(s, cfg['srcn'])

# ================================================================ 14 端到端架构
s = newslide()
page(s, 14, 'Architecture', 'End to end: terminal - network - application',
     'One architecture, in three layers. Each layer is described by what it contains, not by who supplies it.')
LAY = [('TERMINAL', 'What the field carries', '\u7ebf\u7a3f_\u624b\u53f0_2.png',
        [('PNC660 450M handheld', 'B31/B72, IP68, dual SIM'),
         ('MNC650 450M', 'vehicle radio'),
         ('P70 5G+450M', 'radio-style handheld, 2026 Q4'),
         ('PNC660Ex', 'ATEX variant, 2027 Q2')],
        'Hytera 450M Product Introduction PPT_202609_slide9_5.png'),
       ('NETWORK', 'Radio, core and backhaul', '\u901a\u4fe1\u94c1\u5854_2.png',
        [('450 MHz macro eNB', 'wide area coverage, few sites'),
         ('Private core', 'or the operator\u2019s own core'),
         ('Backhaul', 'fibre, microwave or satellite'),
         ('Rapid deployment', 'portable base station and satellite kit')],
        'iBS\u4e00\u4f53\u5316\u57fa\u7ad9_1.png'),
       ('APPLICATION', 'Dispatch, recording and management', '\u7ebf\u7a3f_\u7ec8\u7aef_1.png',
        [('MCS services', 'MCPTT / MCVideo / MCData'),
         ('Dispatch console', 'voice, video, GIS and presence'),
         ('Recording', 'multimedia capture and playback'),
         ('Device management', 'zero-touch deployment')],
        'Hytera PNC660 450MHz&MCX_slide8_3.png')]
for i, (name, sub, ic, items, img) in enumerate(LAY):
    x = X0 + i * 296
    box(s, x, 122, 280, 250, C['CARD'])
    box(s, x, 122, 280, 36, C['NAVY'])
    tx(s, x + 16, 122, 248, 36, name, 13, True, 'WHITE', L, MID)
    box(s, x + 16, 170, 16, 3, C['LIME'])
    fit(s, illu(ic), x + 40, 160, 24, 24)
    tx(s, x + 72, 164, 192, 16, sub, 10, True, 'GREEN')
    for j, (k, d) in enumerate(items):
        yy = 196 + j * 33
        box(s, x + 16, yy + 5, 5, 5, C['LIME'])
        tx(s, x + 28, yy, 240, 14, k, 9.5, True, 'NAVY')
        tx(s, x + 28, yy + 14, 240, 14, d, 8.5, False, 'STEEL')
    box(s, x + 14, 322, 252, 40, C['WHITE'])
    fit(s, illu(img), x + 20, 326, 240, 32)
    if i < 2:
        box(s, x + 282, 232, 12, 26, C['NAVY_L'], shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
box(s, X0, 382, CW, 52, C['PAPER'])
tx(s, X0 + 18, 390, 300, 16, 'THE SAME ARCHITECTURE SERVES ALL SIX INDUSTRIES', 8.5, True, 'GREEN')
for i, t in enumerate(['Power', 'Water', 'Police', 'Emergency', 'Rail', 'Port']):
    x = X0 + 18 + i * 140
    box(s, x, 408, 130, 20, C['CARD'])
    tx(s, x, 408, 130, 20, t, 9, True, 'NAVY', CEN, MID)
concl(s, [('Only the terminals, the applications and the coverage area change.  ', True, 'WHITE'),
          ('The network layer is the same build.', False, 'MIST')])
src(s, '450M Roadmap; HyTalk MC V6.1; 450 MHz Alliance, Equipment Report (May 2025)')

# ================================================================ 15 全球运营商
s = newslide()
page(s, 15, 'Reference  \u00b7  450 MHz', '450 MHz operators worldwide: the model works',
     'Two things are the same in every market: a licensed operator carries the traffic, and utilities anchor the demand.')
box(s, X0, 122, 400, 296, C['WHITE'])
fit(s, illu(cropcache('worldmap_crop.png', asset('alliance_worldmap.png'), (0.12, 0.42, 0.96, 0.72)), 1800),
    X0 + 14, 138, 372, 226)
box(s, X0 + 14, 372, 372, 34, C['CARD'])
tx(s, X0 + 30, 372, 340, 34,
   'Figure 2. The world map of 380 MHz, 410 MHz and 450 MHz deployment.',
   8.5, False, 'STEEL', L, MID, line=1.25)
box(s, X0, 122, 400, 296, None, C['LINE'], 1.0)
OPS = [('Germany', '450connect', 'Band 72, nationwide utility network; more than 95% of Germany covered'),
       ('Austria', 'ArgoNET', 'Band 72; regional CDMA and LTE networks for utilities'),
       ('Poland', 'PGE / Polkomtel', 'Band 31 and band 87; smart grid and field services'),
       ('Malaysia', 'TNB', 'Band 31; procurement for initial deployment in four regions'),
       ('Indonesia', 'SeaTelSat', 'Band 31; rural LTE over satellite backhaul, and a private LTE network for MCX'),
       ('South Africa', 'Transnet', 'Band 31 allocated for rail, with wider use under discussion'),
       ('Singapore', 'Public safety', 'Band 31 allocated for public safety, in service')]
for i2, (c, op, d) in enumerate(OPS):
    y = 122 + i2 * 42
    box(s, 468, y, 444, 38, C['CARD'] if i2 % 2 == 0 else C['PAPER'])
    tx(s, 480, y + 3, 96, 16, c, 9.5, True, 'NAVY')
    tx(s, 480, y + 20, 96, 14, op, 8, True, 'GREEN')
    tx(s, 584, y + 4, 320, 30, d, 8.5, False, 'STEEL', line=1.25)
concl(s, [('The pattern is the same everywhere.  ', True, 'WHITE'),
          ('A licensed operator carries the traffic; utilities and public safety anchor the demand.', False, 'MIST')])
src(s, '450 MHz Alliance, Annual Global Update (Dec 2025) p7-p9, p14; RFI for devices (2025) p5-p7')

# ================================================================ 16-18 MCS 案例
CASES = [
    dict(n=16, kicker='Case 1 of 3  \u00b7  450 MHz licence holder',
         title='Sri Lanka Railways: a national network on public LTE',
         point='Nine lines and 403 stations on a public LTE network. Dialog, the operator, also holds Sri Lanka\u2019s 450 MHz licence.',
         img=asset('Hytera PNC660 450MHz&MCX_slide13_1.jpg'), fy=0.5,
         cap='Sri Lanka Railways - the customer\u2019s own network, photographed on site.',
         stats=[('1,950+', 'MCS terminals delivered across the island network'),
                ('9 / 403', 'railway lines and stations covered'),
                ('450 MHz', 'the operator, Dialog, holds Sri Lanka\u2019s 450 MHz licence')],
         metrics=[('\u221210%', 'train delays'), ('\u221260%', 'management cost'),
                  ('\u221240%', 'track repair time'), ('30 min', 'landslide response')],
         srcn='Hytera showroom briefing (SRC-0134/0135) L211; Hytera PNC660 450MHz&MCX case records; '
              '450 MHz Alliance, Annual Global Update (Dec 2025) p8'),
    dict(n=17, kicker='Case 2 of 3  \u00b7  Railway MCX over 5G',
         title='Hong Kong Light Rail: the first MCS over 5G',
         point='The world\u2019s first railway MCS over 5G - replacing TETRA rather than overlaying it.',
         img=enhance('v6_dispatch.png', '20260917_\u4ea7\u54c1\u6587\u6863_SmartDispatchPlus\u4ea7\u54c1\u4ecb\u7ecd_\u8bfe\u4ef6_\u82f1\u6587_slide5_6.png'),
         fy=0.5,
         cap='Illustrative operations imagery; the customer\u2019s own site photography is not cleared for external use.',
         stats=[('810', 'terminals: 511 handheld, 152 bus-mounted and 147 vehicle'),
                ('5G', 'the world\u2019s first railway MCS running over 5G'),
                ('Full', 'redundancy, including the wireless links between sites')],
         metrics=[('TETRA', 'replaced, not overlaid'), ('Shared', 'operator core network and RAN'),
                  ('Train-borne', 'MCS-compliant dispatch'), ('One', 'network for operations and security')],
         srcn='Hytera showroom briefing L170; Hytera PNC660 450MHz&MCX case records'),
    dict(n=18, kicker='Case 3 of 3  \u00b7  Public safety',
         title='Papua New Guinea: a national command centre',
         point='A police headquarters command centre on Hytera ICC, used for daily operations and for APEC summit security.',
         img=asset('20260903_\u6d77\u672c\u4e8c\u90e8\u57f9\u8bad_HyTalk\u4e0eHyTalk-Pro\u65b9\u6848_\u8bfe\u4ef6_20260905_slide59_3.jpg'),
         fy=0.5,
         cap='Illustrative operations imagery; the customer\u2019s own site photography is not cleared for external use.',
         stats=[('3,000', 'police officers and 600 police vehicles served by one centre'),
                ('APEC', 'summit security: 15,000 foreign visitors and more than 200 meetings'),
                ('360\u00b0', 'city-wide video surveillance, with protection around the clock')],
         metrics=[('60 s', 'police dispatch after the alarm bell rings'),
                  ('VIP', 'automatic vehicle route monitoring'),
                  ('City-wide', 'emergency video communication'),
                  ('24/7', 'protection during the summit')],
         srcn='Hytera case records (Papua New Guinea Police HQ command centre), slides 2-4'),
]
for cfg in CASES:
    s = newslide()
    page(s, cfg['n'], cfg['kicker'], cfg['title'], cfg['point'])
    photo(s, cfg['img'], X0, 122, 470, 246, fy=cfg['fy'])
    box(s, X0, 122, 470, 246, None, C['LINE'], 1.0)
    tx(s, X0, 374, 470, 24, cfg['cap'], 8.5, False, 'SILVER', line=1.3)
    for i, (v, lab) in enumerate(cfg['stats']):
        y = 122 + i * 86
        box(s, 558, y, 354, 76, C['CARD'])
        box(s, 558, y, 3.5, 76, C['NAVY'] if i == 0 else C['MIST'])
        tx(s, 576, y + 10, 140, 28, v, 18, True, 'NAVY')
        tx(s, 576, y + 40, 320, 30, lab, 9.5, False, 'STEEL', line=1.25)
    box(s, X0, 406, CW, 34, C['NAVY'])
    for i, (v, lab) in enumerate(cfg['metrics']):
        x = X0 + 12 + i * 212
        tx(s, x, 409, 204, 18, v, 12, True, 'LIME')
        tx(s, x, 426, 204, 12, lab, 8, False, 'MIST')
    src(s, cfg['srcn'])

# ================================================================ 19 商业模式
s = newslide()
page(s, 19, 'Business model', 'Who does what',
     'Three parties, one chain. The question is not whether the technology works, but who commits first.')
FLOW = [('Hytera', 'supplies the radio, terminals, platform and integration',
         '\u5de5\u5382.png', 'NAVY'),
        ('Globe Telecom', 'holds the spectrum, builds and operates the network, owns the customer relationship',
         '\u901a\u4fe1\u94c1\u5854_3.png', 'GREEN'),
        ('Anchor customers', 'utilities, public safety, transport and ports contract capacity before build-out',
         '\u529e\u516c\u56ed\u533a.png', 'NAVY_L')]
for i, (t, d, ic, col) in enumerate(FLOW):
    x = X0 + i * 300
    box(s, x, 122, 264, 116, C[col])
    box(s, x + 16, 134, 4, 26, C['LIME'])
    tx(s, x + 30, 134, 220, 26, t, 14.5, True, 'WHITE', L, MID)
    tx(s, x + 30, 172, 218, 58, d, 9.5, False, 'MIST', line=1.3)
    if i < 2:
        box(s, x + 268, 166, 26, 34, C['LIME'], shape=MSO_SHAPE.RIGHT_ARROW)
klab(s, X0, 256, 'Three ways to cut the cost')
COST = [('\u901a\u4fe1\u94c1\u5854_2.png', 'Reuse what Globe already has',
         'Towers, fibre, equipment rooms and power are already in place along the routes that matter.'),
        ('\u7ebf\u7a3f_\u624b\u53f0_2.png', 'Far fewer sites for the same coverage',
         'The site count for the same area is a fraction of mid-band - which is what makes backup power affordable.'),
        ('\u529e\u516c\u697c.png', 'Contract the demand before the build',
         'Anchor customers commit capacity first, so the network is built against revenue rather than in front of it.')]
for i, (ic, t, d) in enumerate(COST):
    x = X0 + i * 292
    box(s, x, 280, 280, 100, C['CARD'])
    box(s, x, 280, 3.5, 100, C['LIME'])
    fit(s, illu(ic), x + 16, 294, 28, 40)
    tx(s, x + 54, 292, 210, 22, t, 11, True, 'NAVY', line=1.2)
    tx(s, x + 16, 336, 248, 40, d, 9, False, 'STEEL', line=1.32)
box(s, X0, 392, CW, 44, C['PAPER'])
box(s, X0, 392, 3.5, 44, C['GREEN'])
rich(s, X0 + 18, 392, CW - 36, 44,
     [('A lesson from the Alliance\u2019s country reports:  ', True, 'NAVY'),
      ('Portugal\u2019s 450 MHz trial stalled on commercial terms, not on the technology - so the sequencing matters most.',
       False, 'STEEL')], 11, anchor=MID)
src(s, '450 MHz Alliance, Annual Global Update (Dec 2025) p9 (Portugal: trial network, licence award pending on commercial terms); Hytera proposal')

# ================================================================ 20 收入
s = newslide()
page(s, 20, 'Business model', 'Six revenue streams on one network',
     'A single 450 MHz build supports six commercial lines. None of them depends on consumer take-up.')
REV = [('\u7ebf\u7a3f_\u7ec8\u7aef_1.png', 'IoT connectivity', 'Monthly fee per meter, sensor or asset, on the same sites as the voice service.'),
       ('\u7ebf\u7a3f_\u624b\u53f0_2.png', 'MCX as a service', 'Hosted push-to-talk, video and data, charged per user per month.'),
       ('\u901a\u4fe1\u94c1\u5854_2.png', 'Capacity packages', 'Reserved network share for an anchor customer, on a fixed annual fee.'),
       ('\u624b\u673a.png' if False else '\u7ebf\u7a3f_\u6267\u6cd5\u8bb0\u5f55\u4eea.png', 'Device as a service', 'Terminals bundled with the service, billed monthly rather than bought outright.'),
       ('\u529e\u516c\u697c.png', 'Integration and O&M', 'SCADA, telemetry and smart-city projects, delivered and maintained.'),
       ('\u79fb\u52a8\u65b9\u8231.png', 'Emergency readiness', 'Rapid deployment kit on standby, retained by the province or the agency.')]
for i, (ic, t, d) in enumerate(REV):
    r, c = divmod(i, 3)
    x = X0 + c * 292
    y = 122 + r * 156
    box(s, x, y, 280, 144, C['CARD'])
    box(s, x, y, 280, 3.5, C['NAVY'] if r == 0 else C['LIME'])
    fit(s, illu(ic), x + 18, y + 20, 40, 40)
    box(s, x + 70, y + 30, 16, 3, C['LIME'])
    tx(s, x + 70, y + 40, 194, 36, t, 12.5, True, 'NAVY', line=1.2)
    tx(s, x + 16, y + 78, 248, 56, d, 9.5, False, 'STEEL', line=1.35)
concl(s, [('One build, six lines.  ', True, 'WHITE'),
          ('The utility contract funds the coverage that the other five then use.', False, 'MIST')])
src(s, 'Hytera proposal; 450 MHz Alliance business model material')

# ================================================================ 21 路线图
s = newslide()
page(s, 21, 'Next steps', 'Roadmap: pilot first, then scale',
     'One typhoon-belt province, two anchor customers, and the data Globe needs for a national decision.')
STEPS = [('01', 'Confirm the spectrum', 'Band, bandwidth and licence term', '\u7ebf\u7a3f_\u624b\u53f0_2.png'),
         ('02', 'Pick the pilot region', 'One province, plus two anchors', '\u83f2\u5f8b\u5bbe\u89e3\u51b3\u65b9\u6848\u5305\u88c5_slide11_1.png'),
         ('03', 'Build the pilot', 'Sites, platform, terminals, meters', '\u901a\u4fe1\u94c1\u5854_3.png'),
         ('04', 'Prove the case', 'Coverage, availability and cost data', '\u529e\u516c\u697c.png'),
         ('05', 'Scale nationwide', 'Emergency, transport, smart city', '\u57ce\u5e02\u697c\u7fa4.png')]
for i, (k, t, d, ic) in enumerate(STEPS):
    x = X0 + i * 168
    box(s, x, 122, 186, 96, C['NAVY'] if i == 0 else C['NAVY_L'], shape=MSO_SHAPE.PENTAGON)
    tx(s, x + 14, 132, 150, 16, k, 12, True, 'LIME')
    tx(s, x + 14, 150, 150, 30, t, 10.5, True, 'WHITE', line=1.15)
    tx(s, x + 14, 182, 152, 30, d, 8.5, False, 'MIST', line=1.25)
box(s, X0, 248, CW, 116, C['CARD'])
box(s, X0, 248, 3.5, 116, C['GREEN'])
tx(s, X0 + 20, 262, 300, 20, 'What we need from Globe', 12.5, True, 'NAVY')
NEEDS = [('Spectrum parameters', 'The band and bandwidth, so the coverage and capacity modelling is exact.'),
         ('A candidate province', 'One region, ideally in the typhoon belt, where the gap is visible.'),
         ('Anchor introductions', 'One electric cooperative and one water district willing to trial.')]
for i, (k, d) in enumerate(NEEDS):
    x = X0 + 20 + i * 280
    box(s, x, 296, 6, 6, C['LIME'])
    tx(s, x + 16, 292, 248, 16, k, 10.5, True, 'NAVY')
    tx(s, x + 16, 314, 250, 42, d, 9.5, False, 'STEEL', line=1.3)
box(s, X0, 380, CW, 40, C['GREEN_L'])
rich(s, X0 + 18, 380, CW - 36, 40,
     [('We are not proposing a build.  ', True, 'GREEN'),
      ('The next step is a joint pilot design, so Globe can test the assumptions in its own terrain.',
       False, 'STEEL')], 11, anchor=MID)
concl(s, [('The technical case is settled.  ', True, 'WHITE'),
          ('What decides the outcome is whether demand is contracted before build-out.', False, 'MIST')])
src(s, 'Hytera proposal')

# ================================================================ 22 尾页
s = newslide()
box(s, 0, 0, W, H, C['WHITE'])
box(s, 0, 0, W, 6, C['NAVY'])
box(s, 0, 0, 320, 6, C['LIME'])
box(s, 0, 6, 6, H - 6, C['NAVY'])
fit(s, illu('hytera_logo.png', 900), X1 - 148, 30, 110, 44)

tx(s, X0 + 20, 108, 460, 18, 'HYTERA  \u00d7  GLOBE TELECOM', 10.5, True, 'GREEN')
tx(s, X0 + 20, 138, 470, 60, 'Thank you', 40, True, 'NAVY')
box(s, X0 + 20, 246, 72, 4, C['LIME'])
tx(s, X0 + 20, 272, 440, 22, 'Mission First  \u00b7  Safety Always', 14, True, 'GREEN')
tx(s, X0 + 20, 308, 450, 48,
   'Hytera \u00d7 Globe Telecom  \u00b7  450 MHz for critical infrastructure\nand remote communities.',
   11.5, False, 'STEEL', line=1.42)
tx(s, X0 + 20, 370, 450, 60,
   'Next step: confirm the spectrum parameters and nominate a pilot region,\nand we will design the pilot with you.',
   11.5, False, 'STEEL', line=1.42)

fit(s, illu('\u83f2\u5f8b\u5bbe\u89e3\u51b3\u65b9\u6848\u5305\u88c5_slide11_1.png', 1400), 628, 66, 244, 352)
fit(s, illu('\u7ec8\u7aef_\u591a\u6a21\u624b\u53f0_1.png'), 542, 148, 86, 266)

ln(s, X0, 466, X1, 466, C['LINE'], 1.0)
tx(s, X0, 482, 420, 16, 'Hytera Communications Corporation Limited', 11.5, True, 'NAVY')
tx(s, 560, 484, 352, 14, 'Thank you for your time', 9.5, False, 'SILVER', R)


# ---------------------------------------------------------------- 备注占位符
# 本模板 notesMaster 的 spTree 里没有 body 占位符，python-pptx 建不出备注文本框，
# ppt_notes.py 会报 'NoneType' has no attribute 'text'。保存前补一个。
from lxml import etree
from pptx.oxml.ns import nsdecls

_BODY_PH = (
    '<p:sp %s><p:nvSpPr><p:cNvPr id="99" name="Notes Placeholder"/>'
    '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
    '<p:nvPr><p:ph type="body" sz="quarter" idx="1"/></p:nvPr></p:nvSpPr>'
    '<p:spPr><a:xfrm><a:off x="685800" y="4343400"/><a:ext cx="5486400" cy="4114800"/></a:xfrm></p:spPr>'
    '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>' % nsdecls('p', 'a')
)
for _part in prs.part.package.iter_parts():
    if str(_part.partname).endswith('notesMaster1.xml'):
        _tree = _part._element.find(qn('p:cSld')).find(qn('p:spTree'))
        if _tree.find('.//' + qn('p:ph')) is None:
            _tree.append(etree.fromstring(_BODY_PH))

prs.save(OUT)
print('saved: %s | slides: %d' % (OUT, len(prs.slides._sldIdLst)))
