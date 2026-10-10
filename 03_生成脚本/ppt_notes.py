# -*- coding: utf-8 -*-
"""把逐页讲稿灌进 PPTX 的演讲者备注。

讲稿格式：每页以 "## Slide N" 开头；【提示】开头的手卡行、"---" 分隔线不进备注；
"## 备用问答" 之后的内容整体跳过。

用法：
  python ppt_notes.py <讲稿.md> <方案.pptx> [输出.pptx]       灌备注（缺省原地覆盖）
  python ppt_notes.py --strip <方案.pptx> <对外版.pptx>        清空全部备注，出对外版
"""
import re
import sys
from pathlib import Path

from pptx import Presentation

LIMIT = 520   # 单页备注字符上限，与讲稿约定一致


def parse(md):
    notes, cur = {}, None
    for line in Path(md).read_text(encoding='utf-8').splitlines():
        m = re.match(r'^##\s+Slide\s+(\d+)\s*$', line)
        if m:
            cur = int(m.group(1)); notes[cur] = []
            continue
        if line.startswith('## '):          # 备用问答等非页面章节
            cur = None
            continue
        t = line.strip()
        if cur is None or not t or t == '---' or t.startswith('【提示】'):
            continue
        notes[cur].append(t)
    return {k: '\n'.join(v) for k, v in notes.items()}


def body_frame(slide):
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    if tf is None:
        sys.exit('notesMaster 没有 body 占位符，建不出备注文本框——先用 build 脚本重建（它会补占位符）')
    return tf


def fill(md, pptx, out):
    notes = parse(md)
    prs = Presentation(pptx)
    n = len(prs.slides)
    missing = [i for i in range(1, n + 1) if i not in notes]
    extra = [k for k in notes if k > n]
    if extra:
        sys.exit('讲稿有 Slide %s，但 PPTX 只有 %d 页——页数对不上，先确认构建是否成功' % (extra, n))
    over = []
    for i, s in enumerate(prs.slides, 1):
        text = notes.get(i, '')
        body_frame(s).text = text
        if len(text) > LIMIT:
            over.append((i, len(text)))
    prs.save(out)
    print('notes: %d/%d slides filled -> %s' % (n - len(missing), n, out))
    if missing:
        print('  WARN 无讲稿的页：%s' % missing)
    for i, L in over:
        print('  WARN Slide %d 备注 %d 字符，超过上限 %d' % (i, L, LIMIT))


def strip(pptx, out):
    prs = Presentation(pptx)
    k = 0
    for s in prs.slides:
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            if s.notes_slide.notes_text_frame.text.strip():
                k += 1
            s.notes_slide.notes_text_frame.text = ''
    prs.save(out)
    print('stripped notes on %d slides -> %s' % (k, out))


if __name__ == '__main__':
    a = sys.argv[1:]
    if a[:1] == ['--strip'] and len(a) == 3:
        strip(a[1], a[2])
    elif len(a) in (2, 3):
        fill(a[0], a[1], a[2] if len(a) == 3 else a[1])
    else:
        sys.exit(__doc__)
