# -*- coding: utf-8 -*-
"""把项目里的 Markdown 说明文档转成 Word（中英混排、支持表格）。

用法：python md2docx.py <in.md> <out.docx> [--landscape]
"""
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

CN_FONT = '微软雅黑'
EN_FONT = 'Arial'
NAVY = RGBColor(0x0B, 0x30, 0x8E)
GREEN = RGBColor(0x7F, 0x9B, 0x00)
INK = RGBColor(0x16, 0x21, 0x3E)


def set_font(run, size=10.5, bold=False, color=INK, cn=CN_FONT, en=EN_FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rf)
    rf.set(qn('w:ascii'), en)
    rf.set(qn('w:hAnsi'), en)
    rf.set(qn('w:eastAsia'), cn)


def add_runs(par, text, size=10.5, base_color=INK):
    """处理 **粗体** 与 `代码`。"""
    for part in re.split(r'(\*\*.+?\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            set_font(par.add_run(part[2:-2]), size, True, base_color)
        elif part.startswith('`') and part.endswith('`'):
            set_font(par.add_run(part[1:-1]), size - 0.5, False, RGBColor(0x33, 0x44, 0x66), 'Consolas', 'Consolas')
        else:
            set_font(par.add_run(part), size, False, base_color)


def parse_table(lines):
    rows = []
    for ln in lines:
        cells = [c.strip() for c in ln.strip().strip('|').split('|')]
        if all(re.fullmatch(r':?-{2,}:?', c) for c in cells if c):
            continue
        rows.append(cells)
    return rows


def main():
    src, dst = sys.argv[1], sys.argv[2]
    landscape = '--landscape' in sys.argv
    md = open(src, encoding='utf-8').read().split('\n')

    doc = Document()
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.8)

    i = 0
    while i < len(md):
        ln = md[i]
        s = ln.strip()

        if s.startswith('|') and i + 1 < len(md) and set(md[i + 1].strip()) <= set('|-: '):
            block = []
            while i < len(md) and md[i].strip().startswith('|'):
                block.append(md[i]); i += 1
            rows = parse_table(block)
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=0, cols=ncol)
                t.style = 'Table Grid'
                for ri, r in enumerate(rows):
                    cells = t.add_row().cells
                    for ci in range(ncol):
                        txt = r[ci] if ci < len(r) else ''
                        p = cells[ci].paragraphs[0]
                        p.paragraph_format.space_before = Pt(1)
                        p.paragraph_format.space_after = Pt(1)
                        add_runs(p, txt, 8.5 if landscape else 8,
                                 NAVY if ri == 0 else INK)
                        if ri == 0:
                            for run in p.runs:
                                run.font.bold = True
                doc.add_paragraph()
            continue

        if s.startswith('# '):
            p = doc.add_paragraph(); add_runs(p, s[2:], 17, NAVY)
            for r in p.runs: r.font.bold = True
        elif s.startswith('## '):
            p = doc.add_paragraph(); add_runs(p, s[3:], 13.5, NAVY)
            for r in p.runs: r.font.bold = True
            p.paragraph_format.space_before = Pt(12)
        elif s.startswith('### '):
            p = doc.add_paragraph(); add_runs(p, s[4:], 11.5, GREEN)
            for r in p.runs: r.font.bold = True
            p.paragraph_format.space_before = Pt(8)
        elif s.startswith('> '):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6)
            add_runs(p, s[2:], 9.5, RGBColor(0x55, 0x60, 0x75))
        elif s.startswith('- [ ] '):
            p = doc.add_paragraph(style='List Bullet'); add_runs(p, '☐ ' + s[6:])
        elif s.startswith('- '):
            p = doc.add_paragraph(style='List Bullet'); add_runs(p, s[2:])
        elif s == '---':
            doc.add_paragraph()
        elif s:
            p = doc.add_paragraph(); add_runs(p, s)
        i += 1

    doc.save(dst)
    print('saved:', dst)


if __name__ == '__main__':
    main()
